"""
Formatted XLSX exporters for dispensation reports (GR 3.47 and GR 5.16).
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def _border():
    thin = Side(style='thin')
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _font(bold=False, size=10):
    return Font(bold=bold, size=size, name='Calibri')


def _align(wrap=True, h='left', v='top'):
    return Alignment(wrap_text=wrap, horizontal=h, vertical=v)


def _col_widths(ws):
    for i, w in enumerate([8, 28, 32, 22, 24, 6, 65, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _force_xlsx(path: str) -> str:
    if not path.lower().endswith('.xlsx'):
        base = path.rsplit('.', 1)[0] if '.' in path.split('/')[-1] else path
        return base + '.xlsx'
    return path


def write_disp347_xlsx(dmvt: pd.DataFrame, path_347: str, stn: str) -> None:
    """Write GR 3.47 dispensation report to formatted XLSX."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'DISP 347'
    yellow = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    center   = _align(h='center', v='center')
    wrap_top = _align()
    _col_widths(ws)
    stn_u = stn.upper()

    # ── Header block ─────────────────────────────────────────────────────────
    _hdr_rows = [
        (f'No. ___/____/PCOM/{stn_u}/Disp under GR 3.47(1) [Non isolated] Dated ___________',
         'Annexure C3-47-1 JPO', 1, 20),
        ('WESTERN RAILWAY', None, 2, 24),
        ('APPLICATION FOR DISPENSATION UNDER GR 3.47 (1) SPECIAL INSTRUCTIONS (NON ISOLATED MOVEMENTS)',
         None, 3, 22),
        ('From:', 'To:', 4, 18),
        (f'Sr. DSTE-{stn_u},\nSr. DOM (G) -{stn_u},\nSr. DOM-{stn_u}',
         'Principal Chief Operations Manager,\nWestern Railway – CCG, Mumbai -400 020', 5, 50),
        ('Through: CSTE (Works-1) CCG, Western Railway', None, 6, 18),
        (f'sub: Dispensation under special instructions under GR. 3.47 (1) for {stn_u} STATION '
         f'in connection with the proposed work at {stn_u} STATION.', None, 7, 50),
    ]
    for val_a, val_h, row, ht in _hdr_rows:
        ws.cell(row=row, column=1, value=val_a)
        if val_h:
            ws.cell(row=row, column=8, value=val_h)
            if row in (4, 5):
                ws.merge_cells(f'A{row}:D{row}')
                ws.merge_cells(f'E{row}:H{row}')
                ws.cell(row=row, column=5, value=val_h)
            else:
                ws.merge_cells(f'A{row}:G{row}')
                ws[f'H{row}'].alignment = _align(h='right', v='center')
        else:
            ws.merge_cells(f'A{row}:H{row}')
        ws[f'A{row}'].font = _font(bold=(row in (1, 2, 3, 4)))
        ws[f'A{row}'].alignment = center if row == 2 else (
            center if row == 3 else wrap_top)
        ws.row_dimensions[row].height = ht

    # ── Table column headers ──────────────────────────────────────────────────
    hdrs = [
        'Sr.', 'Name of the\nstation and work',
        'Rules under which\ndispensation is asked for\n(Rules to be reproduced)',
        'Reasons for\nDeviation', 'Deviation\nProposed', '',
        'Special instructions for approval to be included in SWR rules after approval.\n'
        '(Non isolated movements separation distance more then 300M)', 'Remarks',
    ]
    for col, val in enumerate(hdrs, 1):
        c = ws.cell(row=8, column=col, value=val)
        c.font = _font(bold=True); c.alignment = center; c.border = _border()
    ws.merge_cells('F8:G8')
    ws.row_dimensions[8].height = 70
    for col, val in enumerate([1, 2, 3, 4, 5, None, 6, 7], 1):
        c = ws.cell(row=9, column=col, value=val)
        c.font = _font(bold=True); c.alignment = center; c.border = _border()
    ws.merge_cells('F9:G9')
    ws.row_dimensions[9].height = 16

    # ── Boilerplate text ──────────────────────────────────────────────────────
    rules_text = (
        "GR 3.47 - TAKING 'OFF' SIGNALS FOR MORE THAN ONE TRAIN AT A TIME.- "
        "When two or more trains are approaching simultaneously from any direction, "
        "the signals for one train only shall be taken 'OFF', other necessary signals "
        "being kept at 'ON', until the train for which the signals have been taken 'OFF' "
        "has come to a stand at the station, or has cleared the station and the signals "
        "so taken 'OFF' for the said train have been put back to 'ON' except where under "
        "special instructions, the interlocking or the layout of the yard renders a "
        "contrary procedure safe.\n"
        "(1) Taking 'Off' signals for more than one train at a time when two or more "
        "trains are approaching simultaneously from any direction may be permitted over "
        "non isolated lines; under special instructions when requirements of adequate "
        "distance under Rule 3.40 are fulfilled; and under approved special instructions "
        "when requirements of adequate distance under Rule 3.40 are fulfilled.\n"
        "Note: GR 3.47 is to be read along with the instruction given in annexure-1 at "
        "the end of chapter-III."
    )
    reasons_text = (
        "To avoid delay for want of physical isolation, this dispensation is required "
        "for simultaneous movements due to more frequency of trains in UP/DN direction."
    )
    deviation_text = (
        "Dispensation under special instruction is sought under GR 3.47(1) for Taking "
        "Off simultaneous signals for more than one train at a time when two or more "
        "trains are approaching simultaneously from any direction."
    )
    remarks_text = (
        f"The dispensation will be incorporated in SWR of {stn_u} station "
        "under 'special instructions' after approval."
    )

    # ── Data rows ─────────────────────────────────────────────────────────────
    current_row = 10
    mn_serial = 0
    for _, row_data in dmvt.iterrows():
        mn_text = str(row_data.get('MN-MOVT-LIT', '')).strip()
        sub_raw = str(row_data.get('SUB-MOVT-LIT', '')).strip()
        if not mn_text:
            continue
        mn_serial += 1
        sub_movts  = [s.strip() for s in sub_raw.split('\n') if s.strip()]
        merge_start = current_row

        for col, val in enumerate(
                [str(mn_serial), f'Station-\n{stn_u}', rules_text,
                 reasons_text, deviation_text, None, mn_text, remarks_text], 1):
            if val is not None:
                c = ws.cell(row=current_row, column=col, value=val)
                c.border = _border(); c.alignment = wrap_top; c.font = _font()
        ws.cell(row=current_row, column=1).alignment = center
        ws.cell(row=current_row, column=1).font = _font(bold=True)
        ws.cell(row=current_row, column=7).fill = yellow
        ws.cell(row=current_row, column=7).font = _font(bold=True)
        ws.row_dimensions[current_row].height = 45
        current_row += 1

        for sub_idx, sub_text in enumerate(sub_movts, 1):
            ws.cell(row=current_row, column=1, value=f'{mn_serial}({sub_idx})')
            ws.cell(row=current_row, column=6, value=sub_idx)
            ws.cell(row=current_row, column=7, value=sub_text)
            for col in range(1, 9):
                c = ws.cell(row=current_row, column=col)
                c.border = _border(); c.alignment = wrap_top; c.font = _font()
            ws.cell(row=current_row, column=1).alignment = center
            ws.cell(row=current_row, column=6).alignment = center
            ws.row_dimensions[current_row].height = 60
            current_row += 1

        merge_end = current_row - 1
        if merge_end > merge_start:
            for col_letter in ('B', 'C', 'D', 'E', 'H'):
                ws.merge_cells(f'{col_letter}{merge_start}:{col_letter}{merge_end}')
                top = ws[f'{col_letter}{merge_start}']
                top.alignment = wrap_top; top.border = _border()

    current_row += 1
    for col, val in enumerate(
            ['SR. DSTE', None, 'SR. DOM(G)', None, 'SR. DOM', None, 'CSTE', 'PCOM'], 1):
        c = ws.cell(row=current_row, column=col, value=val)
        c.font = _font(bold=True); c.alignment = center; c.border = _border()
    ws.row_dimensions[current_row].height = 30

    wb.save(_force_xlsx(path_347))


def write_disp516_xlsx(dmvt: pd.DataFrame, path_516: str, stn: str) -> None:
    """Write GR 5.16 dispensation report to formatted XLSX."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'DISP 516'
    yellow = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    center   = _align(h='center', v='center')
    wrap_top = _align()
    _col_widths(ws)
    stn_u = stn.upper()

    # ── Header block ─────────────────────────────────────────────────────────
    for row, (val, ht) in enumerate([
        (f'No. ___/____/PCOM/{stn_u}/Disp under GR 5.16 [Non isolated] Dated ___________', 20),
        ('WESTERN RAILWAY', 24),
        ('APPLICATION FOR DISPENSATION UNDER SPECIAL INSTRUCTIONS UNDER GR 5.16', 22),
    ], 1):
        ws.cell(row=row, column=1, value=val)
        ws.merge_cells(f'A{row}:H{row}' if row > 1 else f'A{row}:G{row}')
        ws[f'A{row}'].font = _font(bold=True, size=13 if row == 2 else 10)
        ws[f'A{row}'].alignment = center if row in (2, 3) else wrap_top
        ws.row_dimensions[row].height = ht
    ws.cell(row=1, column=8, value='Annexure D 5-16 JPO')
    ws['H1'].font = _font(bold=True)
    ws['H1'].alignment = _align(h='right', v='center')

    for row, (a_val, h_val, ht) in enumerate([
        ('From:', 'To:', 18),
        (f'Sr. DSTE-{stn_u},\nSr. DOM (G) -{stn_u},\nSr. DOM-{stn_u}',
         'Principal Chief Operations Manager,\nWestern Railway – CCG, Mumbai -400 020', 50),
        ('Through: CSTE (Works-1) CCG, Western Railway', None, 18),
        (f'sub: Dispensation under special instructions under GR. 5.16 for {stn_u} STATION '
         f'in connection with the proposed work at {stn_u} STATION.', None, 50),
    ], 4):
        ws.cell(row=row, column=1, value=a_val)
        if h_val:
            ws.merge_cells(f'A{row}:D{row}'); ws.merge_cells(f'E{row}:H{row}')
            ws.cell(row=row, column=5, value=h_val)
            ws[f'E{row}'].alignment = wrap_top
        else:
            ws.merge_cells(f'A{row}:H{row}')
        ws[f'A{row}'].font = _font(bold=(row == 4))
        ws[f'A{row}'].alignment = wrap_top
        ws.row_dimensions[row].height = ht

    hdrs = [
        'Sr.', 'Name of the\nstation and work',
        'Rules under which\ndispensation is asked for\n(Rules to be reproduced)',
        'Reasons for\nDeviation', 'Deviation\nProposed', '',
        'Special instructions suggested for approval of non-isolated shunting movements '
        'during reception of trains\n(Non isolated movements separation distance more then 300M)',
        'Remarks',
    ]
    for col, val in enumerate(hdrs, 1):
        c = ws.cell(row=8, column=col, value=val)
        c.font = _font(bold=True); c.alignment = center; c.border = _border()
    ws.merge_cells('F8:G8')
    ws.row_dimensions[8].height = 70
    for col, val in enumerate([1, 2, 3, 4, 5, None, 6, 7], 1):
        c = ws.cell(row=9, column=col, value=val)
        c.font = _font(bold=True); c.alignment = center; c.border = _border()
    ws.merge_cells('F9:G9')
    ws.row_dimensions[9].height = 16

    rules_text = (
        "GR 5.16 - Shunting during reception of trains:- "
        "When signals have been taken 'off' for an incoming train on to a line which is not isolated, "
        "no shunting movement shall be carried out towards points over which the incoming train is to pass "
        "except under special instructions for identified stations where frequent shunting movements take place, "
        "and where such points are protected by a Stop Signal or by a Shunt Signal with the precautions to be "
        "observed while performing shunting that-\n"
        "(a) Shunting shall be carried out under supervision of authorized competent railway servant: and\n"
        "(b) Rake or load should be fully on air brake: and\n"
        "(c) The maximum speed during shunting operations shall not exceed 15KMPH.\n"
        "Note: GR 5.16 is to be read along with the instructions given in Annexure – I at the end of chapter 5 of G & SR"
    )
    reasons_text = (
        "To avoid delay in shunting movement for want of physical isolation during reception / "
        "dispatch of train this dispensation is required for shunting movement due to high frequency "
        "of train in UP / DN direction."
    )
    deviation_text = (
        "Dispense on special instruction by authorized officer is sought under SR 5.16 for "
        "shunting during reception of train."
    )
    remarks_text = (
        "1. The Precautions as per GR 5.16 and annexure-I to GR 5.16 will be followed.\n"
        "2. Precautions under GR 5.20 shall also be followed if conditions apply.\n"
        "3. Same shall be included in SWR and SIP after approval of dispensation under GR 5.16."
    )

    current_row = 10
    mn_serial = 0
    for _, row_data in dmvt.iterrows():
        sh_text = str(row_data.get('SH-DISP-LIT', '')).strip()
        if not sh_text:
            continue
        mn_serial += 1
        split_marker = '::-\n'
        if split_marker in sh_text:
            parts = sh_text.split(split_marker, 1)
            mn_header = parts[0] + '::-'
            sub_part = parts[1].rstrip('.')
        else:
            mn_header = sh_text
            sub_part = ''
        sub_movts = [s.strip() for s in sub_part.split('\n') if s.strip()]
        merge_start = current_row

        for col, val in enumerate(
                [str(mn_serial), f'Station-\n{stn_u}', rules_text,
                 reasons_text, deviation_text, mn_serial, mn_header, remarks_text], 1):
            c = ws.cell(row=current_row, column=col, value=val)
            c.border = _border(); c.alignment = wrap_top; c.font = _font()
        ws.cell(row=current_row, column=1).alignment = center
        ws.cell(row=current_row, column=1).font = _font(bold=True)
        ws.cell(row=current_row, column=6).alignment = center
        ws.cell(row=current_row, column=7).fill = yellow
        ws.cell(row=current_row, column=7).font = _font(bold=True)
        ws.row_dimensions[current_row].height = 45
        current_row += 1

        for sub_idx, sub_text in enumerate(sub_movts, 1):
            ws.cell(row=current_row, column=6, value=sub_idx)
            ws.cell(row=current_row, column=7, value=sub_text)
            for col in range(1, 9):
                c = ws.cell(row=current_row, column=col)
                c.border = _border(); c.alignment = wrap_top; c.font = _font()
            ws.cell(row=current_row, column=6).alignment = center
            ws.row_dimensions[current_row].height = 60
            current_row += 1

        merge_end = current_row - 1
        if merge_end > merge_start:
            for col_letter in ('B', 'C', 'D', 'E', 'H'):
                ws.merge_cells(f'{col_letter}{merge_start}:{col_letter}{merge_end}')
                top = ws[f'{col_letter}{merge_start}']
                top.alignment = wrap_top; top.border = _border()

    current_row += 1
    cert_text = (
        f'Reference Signaling Plan No. SG(SP)/___/___/{stn_u} & TOC No. SG(SG)/___/___/{stn_u} '
        f'of {stn_u} STATION.\n\n'
        'We, hereby certify the above shunting movements during reception of train are permitted '
        'under GR. 5.16.\nRecommended and Forwarded by:- CSTE/W-1/CCG WESTERN RAILWAY'
    )
    ws.cell(row=current_row, column=1, value=cert_text)
    ws.merge_cells(f'A{current_row}:H{current_row}')
    ws[f'A{current_row}'].font = _font()
    ws[f'A{current_row}'].alignment = wrap_top
    ws.row_dimensions[current_row].height = 70
    current_row += 1

    for col, val in enumerate(
            ['SR. DSTE', None, 'SR. DOM(G)', None, 'SR. DOM', None, 'CSTE', 'PCOM'], 1):
        c = ws.cell(row=current_row, column=col, value=val)
        c.font = _font(bold=True); c.alignment = center; c.border = _border()
    ws.row_dimensions[current_row].height = 30

    wb.save(_force_xlsx(path_516))
